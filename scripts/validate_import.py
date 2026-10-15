"""
validate_import.py — بيقرا القوالب المملوءة بنفس طريقة المنظومة ويعمل فحوصات السلامة قبل الاستيراد.
التشغيل:  python -I scripts/validate_import.py [data-import/filled-templates] [templates] [data-import/validation_report.md]

طريقة القراية (زي الـ parser في المنظومة):
  - نوع القالب من الشيت المخفي _meta (B1 = template_id).
  - كل شيت بالاسم. شيت الجدول: الصف 2 = أسماء الأعمدة التقنية، البيانات من الصف 3،
    والصف اللي أوله "مثال" أو فاضي بيتجاهل. الإلزامي = عنوان الصف 1 آخره "*".
  - شيت مفتاح/قيمة (رأس العقد/أمر التكليف): العمود C = المفتاح، B = القيمة، والإلزامي عنوانه آخره "*".
  - قوايم الاختيار والأرقام من الـ Data Validation اللي في الملف.
الدرجات: خطأ يمنع الاستيراد / تحذير / ملاحظة.
"""
import glob
import os
import re
import sys
from collections import Counter, OrderedDict, defaultdict

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
FILLED = sys.argv[1] if len(sys.argv) > 1 else os.path.join("data-import", "filled-templates")
TPL = sys.argv[2] if len(sys.argv) > 2 else "templates"
REPORT = sys.argv[3] if len(sys.argv) > 3 else os.path.join("data-import", "validation_report.md")
EXTRACT_REPORT = os.path.join(os.path.dirname(os.path.abspath(FILLED)), "_report.txt")

VAT = 1.14
VAT_TOL = 0.005
SKIP_SHEETS = {"تعليمات", "_meta"}

ERR, WARN, INFO = "error", "warn", "info"
SEV_TITLE = {ERR: "خطأ يمنع الاستيراد", WARN: "تحذير", INFO: "ملاحظة"}
issues = []  # dict(sev, cat, file, sheet, row, msg)


def add(sev, cat, file="", sheet="", row="", msg=""):
    issues.append(dict(sev=sev, cat=cat, file=file, sheet=sheet, row=row, msg=msg))


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ───────────────────────── القراية ─────────────────────────
def dv_rules(ws):
    """{column_letter: ('list', [..]) | ('number', None) | ('date', None)}"""
    rules = {}
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            letter = re.match(r"[A-Z]+", rng).group(0)
            if dv.type == "list" and dv.formula1 and dv.formula1.startswith('"'):
                rules[letter] = ("list", dv.formula1.strip('"').split(","))
            elif dv.type == "decimal":
                rules[letter] = ("number", None)
            elif dv.type == "date":
                rules[letter] = ("date", None)
    return rules


def read_table(ws):
    keys = [s(c.value) for c in ws[2]]
    labels = [s(c.value) for c in ws[1]]
    required = {k for k, l in zip(keys, labels) if k and l.endswith("*")}
    rules = {keys[column_index_from_string(L) - 1]: r for L, r in dv_rules(ws).items()
             if column_index_from_string(L) - 1 < len(keys)}
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if all(blank(v) for v in r) or s(r[0]).startswith("مثال"):
            continue
        d = {k: v for k, v in zip(keys, r) if k}
        rows.append(d)
    # رقم الصف في الإكسل
    excel_rows = [i for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), 3)
                  if not (all(blank(v) for v in r) or s(r[0]).startswith("مثال"))]
    for d, n in zip(rows, excel_rows):
        d["_row"] = n
    return dict(kind="table", keys=keys, required=required, rules=rules, rows=rows)


def read_kv(ws):
    values, required, rownum = {}, set(), {}
    for r in range(2, ws.max_row + 1):
        key = s(ws.cell(row=r, column=3).value)
        if not key:
            continue
        values[key] = ws.cell(row=r, column=2).value
        rownum[key] = r
        if s(ws.cell(row=r, column=1).value).endswith("*"):
            required.add(key)
    return dict(kind="kv", values=values, required=required, rownum=rownum)


def read_workbook(path):
    wb = load_workbook(path, data_only=True)
    tid = s(wb["_meta"]["B1"].value) if "_meta" in wb.sheetnames else ""
    sheets = OrderedDict()
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        sheets[ws.title] = read_kv(ws) if s(ws.cell(row=1, column=3).value) == "(المفتاح)" else read_table(ws)
    return dict(path=path, name=os.path.basename(path), template_id=tid, sheets=sheets)


official_lists = {}  # (template_id, sheet, key) -> قايمة القالب الرسمي


def load_official(template_id):
    p = os.path.join(TPL, f"{template_id}.xlsx")
    if not os.path.exists(p) or template_id in {k[0] for k in official_lists}:
        return
    wb = load_workbook(p)
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS or s(ws.cell(row=1, column=3).value) == "(المفتاح)":
            continue
        keys = [s(c.value) for c in ws[2]]
        for L, (kind, vals) in dv_rules(ws).items():
            i = column_index_from_string(L) - 1
            if kind == "list" and i < len(keys):
                official_lists[(template_id, ws.title, keys[i])] = vals
    official_lists.setdefault((template_id, "", ""), [])


# ───────────────────────── فحوصات عامة لكل شيت ─────────────────────────
def generic_checks(book):
    off_viol = defaultdict(Counter)
    for sheet, sh in book["sheets"].items():
        if sh["kind"] == "kv":
            for k in sh["required"]:
                if blank(sh["values"].get(k)):
                    add(ERR, "حقل إلزامي فاضي", book["name"], sheet, sh["rownum"][k], f"`{k}` فاضي")
            continue
        for d in sh["rows"]:
            for k in sh["required"]:
                if blank(d.get(k)):
                    add(ERR, "حقل إلزامي فاضي", book["name"], sheet, d["_row"], f"`{k}` فاضي")
            for k, (kind, vals) in sh["rules"].items():
                v = d.get(k)
                if blank(v):
                    continue
                if kind == "list" and s(v) not in vals:
                    add(ERR, "قيمة مش من قائمة الاختيار", book["name"], sheet, d["_row"], f"`{k}` = {s(v)}")
                elif kind == "number" and (not isinstance(v, (int, float)) or v < 0):
                    add(ERR, "قيمة رقمية غلط", book["name"], sheet, d["_row"], f"`{k}` = {s(v)}")
                if kind == "list":
                    off = official_lists.get((book["template_id"], sheet, k))
                    if off is not None and s(v) not in off:
                        off_viol[(sheet, k)][s(v)] += 1
    for (sheet, k), cnt in off_viol.items():
        add(WARN, "قيمة مش موجودة في قائمة القالب الرسمي", book["name"], sheet, "",
            f"`{k}`: " + "، ".join(f"«{v}» ×{n}" for v, n in cnt.items())
            + " — النسخة المملوءة فيها القيم دي، لازم تتضاف لقائمة القالب في scripts/build_templates.py وللـ enum في المنظومة")


# ───────────────────────── التشغيل ─────────────────────────
books = [read_workbook(p) for p in sorted(glob.glob(os.path.join(FILLED, "*.xlsx"))) if not os.path.basename(p).startswith("~$")]
if not books:
    sys.exit(f"مفيش ملفات في {FILLED}")
for b in books:
    load_official(b["template_id"])
    generic_checks(b)

by_tid = defaultdict(list)
for b in books:
    by_tid[b["template_id"]].append(b)

# ── 03 الشركات
companies = {}
for b in by_tid["03_الشركات"]:
    for d in b["sheets"]["الشركات"]["rows"]:
        code = s(d.get("code"))
        if code in companies:
            add(ERR, "كود شركة مكرر", b["name"], "الشركات", d["_row"], code)
        companies[code] = d

# ── 05 العقود والمقايسة
contracts = {}           # contract_no -> header values
boq = {}                 # contract_no -> {item_code: row}
for b in by_tid["05_العقد_والمقايسة"]:
    hdr = b["sheets"]["بيانات العقد"]["values"]
    no = s(hdr.get("contract_no"))
    if no in contracts:
        add(ERR, "العقد مرفوع أكتر من مرة", b["name"], "بيانات العقد", "", no)
    contracts[no] = dict(hdr, _file=b["name"])
    if companies and s(hdr.get("company_code")) not in companies:
        add(ERR, "كود الشركة مش معرّف في قالب الشركات", b["name"], "بيانات العقد", "", s(hdr.get("company_code")))
    for k in ("retention_pct", "tax_pct", "stamp_pct", "advance_pct", "award_date", "start_date", "end_date"):
        if blank(hdr.get(k)):
            add(INFO, "بيانات عقد اختيارية ناقصة (مطلوبة للاستقطاعات/الـ SLA)", b["name"], "بيانات العقد", "", f"{no}: `{k}`")
    items, seen = {}, Counter()
    rows = b["sheets"]["المقايسة"]["rows"]
    for d in rows:
        code = s(d.get("item_code")).replace(".", "-")
        seen[code] += 1
        items[code] = d
    for code, n in seen.items():
        if n > 1:
            add(ERR, "رقم بند مكرر في نفس العقد", b["name"], "المقايسة", items[code]["_row"], f"{no}: {code} ×{n}")
    leaf_no_price = []
    for code, d in items.items():
        is_group = blank(d.get("unit"))
        d["_group"] = is_group
        parent = "-".join(code.split("-")[:-1])
        if parent and parent not in items:
            add(WARN, "البند الأب مش موجود في المقايسة", b["name"], "المقايسة", d["_row"], f"{no}: {code} (الأب {parent})")
        if is_group and not blank(d.get("unit_price")) and d.get("unit_price"):
            add(WARN, "صف عنوان (من غير وحدة) عليه سعر", b["name"], "المقايسة", d["_row"], f"{no}: {code}")
        if not is_group and not d.get("unit_price"):
            leaf_no_price.append(code)
    if leaf_no_price:
        add(INFO, "بنود في المقايسة من غير سعر (أي فوترة عليها = C5)", b["name"], "المقايسة", "",
            f"{no}: {len(leaf_no_price)} بند — " + "، ".join(leaf_no_price[:40]) + (" …" if len(leaf_no_price) > 40 else ""))
    boq[no] = items

# ── 06 أوامر التكليف
ao_defined = {}  # (contract_no, ao_no) -> [revisions]
ao_summary = []
for b in by_tid["06_أمر_التكليف"]:
    hdr = b["sheets"]["بيانات أمر التكليف"]["values"]
    no, ao, rev = s(hdr.get("contract_no")), s(hdr.get("ao_no")).replace(" ", "").upper(), s(hdr.get("revision")) or "0"
    if re.fullmatch(r"\d+", ao):
        ao = f"PO{int(ao):02d}"
    ao_defined.setdefault((no, ao), []).append(rev)
    rown = b["sheets"]["بيانات أمر التكليف"]["rownum"]
    pj = s(hdr.get("project_code"))
    if "?" in pj:
        add(ERR, "أمر التكليف مش مربوط بمشروع (كود مؤقت)", b["name"], "بيانات أمر التكليف", rown.get("project_code", ""),
            f"{ao} R{rev}: project_code = {pj} — الأدمن لازم يختار المشروع من قالب 02")
    if no not in contracts:
        add(ERR, "أمر التكليف على عقد مش مرفوع", b["name"], "بيانات أمر التكليف", "", no)
        continue
    items = boq[no]
    lines = b["sheets"]["بنود أمر التكليف"]["rows"]
    seen = Counter()
    value, no_price = 0.0, OrderedDict()
    for d in lines:
        code = s(d.get("item_code")).replace(".", "-")
        seen[(code, s(d.get("site_code")))] += 1
        it = items.get(code)
        if it is None:
            add(ERR, "بند في أمر التكليف مش موجود في مقايسة العقد", b["name"], "بنود أمر التكليف", d["_row"], f"{ao}: {code}")
            continue
        if it["_group"]:
            add(ERR, "كمية على صف عنوان (مجموعة) مش بند", b["name"], "بنود أمر التكليف", d["_row"], f"{ao}: {code}")
            continue
        q = d.get("planned_qty") or 0
        if it.get("unit_price"):
            value += q * it["unit_price"]
        else:
            no_price.setdefault(code, 0)
            no_price[code] += q
    for (code, site), n in seen.items():
        if n > 1:
            add(ERR, "نفس البند والموقع مكرر في أمر التكليف", b["name"], "بنود أمر التكليف", "", f"{ao}: {code} {site} ×{n}")
    for code, q in no_price.items():
        add(WARN, "بند في أمر التكليف مالوش سعر في العقد (هيتفوتر C5)", b["name"], "بنود أمر التكليف", "",
            f"{ao} R{rev}: {code} — الكمية المخططة {q:g} — {s(items[code].get('description'))[:70]}")
    sites = {s(d.get("site_code")) for d in lines if not blank(d.get("site_code"))}
    ao_summary.append((b["name"], no, ao, rev, len(lines), len(sites), value, len(no_price)))

# ── 07 المهمات
mat_codes = Counter()
mat_rows = defaultdict(list)
units = Counter()
for b in by_tid["07_كتالوج_المهمات"]:
    for d in b["sheets"]["المهمات"]["rows"]:
        c = s(d.get("material_code"))
        mat_codes[c] += 1
        mat_rows[c].append((b["name"], d["_row"], s(d.get("description"))[:60]))
        units[s(d.get("unit"))] += 1
for c, n in mat_codes.items():
    if n > 1:
        for f, r, desc in mat_rows[c]:
            add(ERR, "رقم شطب مكرر في كتالوج المهمات", f, "المهمات", r, f"{c}: {desc}")
if units:
    up = {u.upper().rstrip("S") for u in units}
    if len(up) < len(units):
        add(INFO, "وحدات مكتوبة بأكتر من شكل في كتالوج المهمات (يفضل توحيدها)", "", "المهمات", "",
            "، ".join(f"{u} ×{n}" for u, n in units.most_common()))

# ── 11 المستخلصات السابقة
inv_by_missing_ao = defaultdict(list)
inv_keys, inv_no_keys = Counter(), defaultdict(set)
ratio_bad = []
unknown_steps = 0
for b in by_tid["11_المستخلصات_السابقة"]:
    for d in b["sheets"]["المستخلصات"]["rows"]:
        no, ao, inv = s(d.get("contract_no")), s(d.get("ao_no")), s(d.get("invoice_no"))
        if no not in contracts:
            add(ERR, "مستخلص على عقد مش مرفوع", b["name"], "المستخلصات", d["_row"], no)
        if (no, ao) not in ao_defined:
            inv_by_missing_ao[(no, ao)].append(inv)
        inv_keys[(no, ao, inv)] += 1
        inv_no_keys[(no, inv)].add(ao)
        ex, inc = d.get("amount_excl_vat"), d.get("amount_incl_vat")
        if isinstance(ex, (int, float)) and isinstance(inc, (int, float)) and ex:
            r = inc / ex
            if abs(r / VAT - 1) > VAT_TOL:
                ratio_bad.append((b["name"], d["_row"], no, ao, inv, ex, inc, r))
        elif isinstance(ex, (int, float)) and blank(inc):
            add(INFO, "مستخلص من غير المبلغ الشامل (السيستم هيحسبه ×1.14)", b["name"], "المستخلصات", d["_row"], f"{no} {ao} #{inv}")
        if blank(d.get("current_step")):
            unknown_steps += 1
for (no, ao, inv), n in inv_keys.items():
    if n > 1:
        add(ERR, "مستخلص مكرر (نفس العقد وأمر التكليف والرقم)", "", "المستخلصات", "", f"{no} {ao} #{inv} ×{n}")
for (no, inv), aos_ in inv_no_keys.items():
    if len(aos_) > 1:
        add(WARN, "نفس رقم المستخلص على أكتر من أمر تكليف في نفس العقد", "", "المستخلصات", "", f"{no} #{inv}: {', '.join(sorted(aos_))}")
for f, r, no, ao, inv, ex, inc, ratio in ratio_bad:
    add(WARN, "نسبة الشامل/غير الشامل ≠ 1.14 (±0.5%)", f, "المستخلصات", r,
        f"{no} {ao} #{inv}: {ex:,.2f} ← {inc:,.2f} (النسبة {ratio:.4f})")

# ── ملاحظات الاستخراج (_report.txt)
if os.path.exists(EXTRACT_REPORT):
    for line in open(EXTRACT_REPORT, encoding="utf-8"):
        if line.startswith("!!"):
            add(WARN, "ملاحظة من استخراج الملفات الأصلية", "_report.txt", "", "", line[2:].strip())
        elif line.startswith("حالات محتاجة مراجعة"):
            add(INFO, "ملاحظة من استخراج الملفات الأصلية", "_report.txt", "", "", line.strip())

# ───────────────────────── التقرير ─────────────────────────
cnt = Counter(i["sev"] for i in issues)
L = ["# تقرير فحص بيانات الاستيراد", "",
     f"الملفات اللي اتفحصت ({len(books)}) من `{FILLED}`:", ""]
for b in books:
    rows = sum(len(sh["rows"]) for sh in b["sheets"].values() if sh["kind"] == "table")
    L.append(f"- `{b['name']}` — قالب `{b['template_id']}` — {rows} صف")
L += ["", "## الملخص", "", "| الدرجة | العدد |", "|---|---|"]
for sev in (ERR, WARN, INFO):
    L.append(f"| {SEV_TITLE[sev]} | {cnt[sev]} |")
L += ["", "### أرقام أساسية", "",
      f"- العقود: {len(contracts)} — بنود المقايسة: " + "، ".join(f"{k}: {len(v)} صف" for k, v in boq.items()),
      f"- أوامر التكليف المعرّفة بكميات: {len(ao_summary)} — المستخلصات المرحّلة: {sum(inv_keys.values())}"
      f" (محطتهم غير محددة: {unknown_steps})",
      f"- أوامر تكليف متذكرة في المستخلصات ومش معرّفة: {len(inv_by_missing_ao)} أمر تكليف عليهم "
      f"{sum(len(v) for v in inv_by_missing_ao.values())} مستخلص",
      f"- أصناف المهمات: {sum(mat_codes.values())} (أكواد مكررة: {sum(1 for n in mat_codes.values() if n > 1)})",
      ""]
if ao_summary:
    L += ["### أوامر التكليف", "", "| الملف | العقد | أمر التكليف | الملحق | سطور | مواقع | القيمة المحسوبة (بأسعار العقد) | بنود بدون سعر |", "|---|---|---|---|---|---|---|---|"]
    for f, no, ao, rev, n, ns, v, np_ in ao_summary:
        L.append(f"| {f} | {no} | {ao} | {rev} | {n} | {ns} | {v:,.2f} | {np_} |")
    L.append("")


def md(v):
    return s(v).replace("|", "\\|").replace("\n", " ")


MAX_ROWS = 80
for sev in (ERR, WARN, INFO):
    sel = [i for i in issues if i["sev"] == sev]
    L += [f"## {SEV_TITLE[sev]} ({len(sel)})", ""]
    if sev == ERR and inv_by_missing_ao:
        L += ["### مستخلصات على أوامر تكليف مش معرّفة "
              f"({sum(len(v) for v in inv_by_missing_ao.values())} مستخلص على {len(inv_by_missing_ao)} أمر تكليف)", "",
              "لازم أمر التكليف يترفع (قالب 06 — على الأقل بيانات الرأس + المشروع) قبل ترحيل مستخلصاته.", "",
              "| العقد | أمر التكليف | عدد المستخلصات | أرقام المستخلصات |", "|---|---|---|---|"]
        for (no, ao), invs in sorted(inv_by_missing_ao.items()):
            L.append(f"| {no} | {ao} | {len(invs)} | {md('، '.join(invs))} |")
        L.append("")
    if not sel and not (sev == ERR and inv_by_missing_ao):
        L += ["لا يوجد.", ""]
        continue
    cats = OrderedDict()
    for i in sel:
        cats.setdefault(i["cat"], []).append(i)
    for cat, lst in cats.items():
        L += [f"### {cat} ({len(lst)})", "", "| الملف | الشيت | الصف | التفاصيل |", "|---|---|---|---|"]
        for i in lst[:MAX_ROWS]:
            L.append(f"| {md(i['file'])} | {md(i['sheet'])} | {md(i['row'])} | {md(i['msg'])} |")
        if len(lst) > MAX_ROWS:
            L.append(f"| … | | | و {len(lst) - MAX_ROWS} كمان |")
        L.append("")

# عدد مستخلصات أوامر التكليف غير المعرفة بيتحسب ضمن الأخطاء
total_err = cnt[ERR] + sum(len(v) for v in inv_by_missing_ao.values())
L[L.index("| خطأ يمنع الاستيراد | %d |" % cnt[ERR])] = f"| {SEV_TITLE[ERR]} | {total_err} |"
os.makedirs(os.path.dirname(os.path.abspath(REPORT)), exist_ok=True)
with open(REPORT, "w", encoding="utf-8") as f:
    f.write("\n".join(L) + "\n")
print(f"{SEV_TITLE[ERR]}: {total_err} | {SEV_TITLE[WARN]}: {cnt[WARN]} | {SEV_TITLE[INFO]}: {cnt[INFO]}")
print("التقرير:", REPORT)
