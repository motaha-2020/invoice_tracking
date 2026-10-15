"""
fill_templates_from_example.py — يملى نسخ من القوالب الرسمية (templates/) بالبيانات الحقيقية المستخرجة في data-import/*.csv
التشغيل:  python -I scripts/fill_templates_from_example.py [templates] [data-import]
الناتج:   data-import/filled-templates/*.xlsx  (نفس شكل القالب بالظبط — جاهزة للرفع على المنظومة)

لكل قالب: بنفتح الملف الرسمي، نمسح صفوف المثال الرمادي، ونكتب الصفوف الحقيقية من الصف 3،
ونسيب الـ Data Validation والتنسيق والتعليقات وشيت التعليمات و_meta زي ما هم.
"""
import csv
import os
import sys
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.stdout.reconfigure(encoding="utf-8")
TPL = sys.argv[1] if len(sys.argv) > 1 else "templates"
SRC = sys.argv[2] if len(sys.argv) > 2 else "data-import"
OUT = os.path.join(SRC, "filled-templates")
os.makedirs(OUT, exist_ok=True)

FONT = "Cairo"
THIN = Side(style="thin", color="DEDEE4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GROUP_FILL = PatternFill("solid", fgColor="EDE5F4")
WARN_FILL = PatternFill("solid", fgColor="FFF4CC")
NO_FILL = PatternFill(fill_type=None)

COMPANY_CODE = {"co-bridge": "CO-BRIDGE"}

# الباب الرئيسي في المقايسة ← نوع العمل.
# 4 و5 بنفس قيم الديمو (ديمو/js/data.js)؛ 6 "أعمال إعادة الشيء لأصله" = الملف الأصلي بيجمّعها تحت "اجمالي اعمال الصيانة".
WORK_TYPE_BY_SECTION = {"1": "أعمال مدنية", "2": "تركيبات", "3": "تخطيط ورسومات",
                        "4": "صواعد داخل المباني", "5": "رفع كفاءة", "6": "صيانة"}
EXTRA_WORK_TYPES = ["صواعد داخل المباني", "رفع كفاءة"]  # مش موجودين في قائمة القالب الرسمي (بنضيفهم في النسخة المملوءة)

# mapped_step (من extract_example_data.py) ← قيمة "المحطة الحالية" في قالب 11
STEP_AR = {"done": "تم الصرف", "stores": "المخازن والمشتريات", "tech": "الشئون الفنية", "gm": "المدير العام",
           "exec": "مدير التنفيذ", "sup": "مشرف التنفيذ", "company": "الشركة"}


def read_csv(name):
    with open(os.path.join(SRC, name), encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def num(v):
    if v in (None, ""):
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def open_template(template_id):
    return load_workbook(os.path.join(TPL, f"{template_id}.xlsx"))


def clear_examples(ws):
    """يمسح صفوف المثال (من الصف 3) — القيم والتلوين — ويسيب تنسيق الأرقام/التواريخ"""
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for c in row:
            if c.value is not None or c.fill.fill_type:
                c.value = None
                c.fill = NO_FILL
                c.font = Font(name=FONT, size=10)
                c.border = Border()


def write_rows(ws, rows, group_flags=None, money_cols=()):
    """rows = list of dict بالمفاتيح التقنية (صف 2)"""
    keys = [c.value for c in ws[2]]
    for i, r in enumerate(rows):
        is_group = bool(group_flags and group_flags[i])
        for j, k in enumerate(keys, 1):
            v = r.get(k)
            c = ws.cell(row=3 + i, column=j, value=None if v == "" else v)
            c.font = Font(name=FONT, size=10, bold=is_group)
            c.border = BORDER
            if is_group:
                c.fill = GROUP_FILL
            if k in money_cols and isinstance(v, (int, float)):
                c.number_format = "#,##0.00"
            if k == "description":
                c.alignment = Alignment(wrap_text=False, vertical="top")
    return len(rows)


def fill_kv(ws, values, notes=None):
    """شيت مفتاح/قيمة: العمود C = المفتاح، B = القيمة، D = ملاحظات"""
    notes = notes or {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=3).value
        if not key:
            continue
        v = values.get(key, "")
        cell = ws.cell(row=r, column=2, value=None if v in ("", None) else v)
        cell.font = Font(name=FONT, bold=False)
        cell.fill = NO_FILL
        if key in notes:
            cell.fill = WARN_FILL
            n = ws.cell(row=r, column=4, value=notes[key])
            n.font = Font(name=FONT, size=9, color="9C5700", bold=True)


def save(wb, name):
    path = os.path.join(OUT, name)
    wb.save(path)
    print(f"  -> {name}")
    return path


# ─────────────── البيانات ───────────────
companies = read_csv("companies.csv")
contracts = read_csv("contracts.csv")
items = read_csv("contract_items.csv")
aos = read_csv("assignment_orders.csv")
ao_items = read_csv("ao_items.csv")
materials = read_csv("materials_catalog.csv")
invoices = read_csv("invoices_history.csv")

items_by_contract = defaultdict(list)
for it in items:
    items_by_contract[it["contract_no"]].append(it)
desc_of = {(it["contract_no"], it["item_code"]): it["description"] for it in items}

MISSING = "ناقص: "
print("ملء القوالب في", OUT)

# ═════════════ 03 الشركات ═════════════
wb = open_template("03_الشركات")
ws = wb["الشركات"]
clear_examples(ws)
write_rows(ws, [dict(code=COMPANY_CODE.get(c["id"], c["id"].upper()), name=c["name"], full_name=c["full_name"]) for c in companies])
save(wb, "03_الشركات.xlsx")

# ═════════════ 05 العقد والمقايسة (ملف لكل عقد) ═════════════
for k in contracts:
    no = k["contract_no"]
    wb = open_template("05_العقد_والمقايسة")
    fill_kv(wb["بيانات العقد"],
            dict(contract_no=no, company_code=COMPANY_CODE.get(k["company_id"], k["company_id"]), title=k["title"],
                 currency=k["currency"], vat_pct=num(k["vat_pct"])),
            notes={"award_date": MISSING + "تاريخ برقية القبول مش موجود في ملفات الإكسل (موجود في PDF البرقية)",
                   "start_date": MISSING + "من برقية القبول / محضر التسليم",
                   "retention_pct": MISSING + "نسبة تأمين ضمان الأعمال — من شروط العقد",
                   "tax_pct": MISSING + "نسبة الخصم من المنبع — من المالية",
                   "stamp_pct": MISSING + "نسبة الدمغة — من المالية",
                   "advance_pct": MISSING + "الدفعة المقدمة (لو فيه) — من شروط العقد"})
    ws = wb["المقايسة"]
    clear_examples(ws)
    # نضيف قيم نوع العمل الناقصة في قائمة الاختيار (النسخة المملوءة بس)
    for dv in ws.data_validations.dataValidation:
        if "F3" in str(dv.sqref) and dv.formula1 and dv.formula1.startswith('"'):
            vals = dv.formula1.strip('"').split(",")
            dv.formula1 = '"' + ",".join(vals + [v for v in EXTRA_WORK_TYPES if v not in vals]) + '"'
    rows, flags = [], []
    for it in items_by_contract[no]:
        g = it["is_group"] == "1"
        top = it["item_code"].split("-")[0]
        rows.append(dict(item_code=it["item_code"], description=it["description"], unit=it["unit"],
                         contract_qty=None if g else num(it["contract_qty"]), unit_price=None if g else num(it["unit_price"]),
                         work_type="" if g else WORK_TYPE_BY_SECTION.get(top, "")))
        flags.append(g)
    n = write_rows(ws, rows, flags, money_cols=("unit_price",))
    save(wb, f"05_العقد_والمقايسة__{no.replace('/', '-')}.xlsx")
    print(f"     {n} صف ({sum(flags)} عنوان، {n - sum(flags)} بند)")

# ═════════════ 06 أوامر التكليف (ملف لكل أمر تكليف بكميات) ═════════════
ao_meta = {(a["contract_no"], a["ao_no"]): a for a in aos}
lines_by_ao = defaultdict(list)
for a in ao_items:
    lines_by_ao[(a["contract_no"], a["ao_no"])].append(a)
for (contract_no, ao_key), lines in lines_by_ao.items():
    meta = ao_meta.get((contract_no, ao_key), {})
    base_no, _, rev = ao_key.partition("-R")
    wb = open_template("06_أمر_التكليف")
    fill_kv(wb["بيانات أمر التكليف"],
            dict(contract_no=contract_no, ao_no=base_no, revision=int(rev) if rev else 0,
                 project_code="PJ-?", name=meta.get("name", "")),
            notes={"project_code": "لازم الأدمن يربطه: PJ-? مكان مؤقت — اختار كود المشروع من قالب 02 المشروعات "
                                   "(ملفات الإكسل مفيهاش ربط أمر التكليف بالمشروع/السنترال)",
                   "ao_date": MISSING + "تاريخ أمر التكليف موجود في PDF أمر التكليف بس",
                   "ao_value": "فاضي = السيستم يحسبها من الكميات × أسعار العقد"})
    ws = wb["بنود أمر التكليف"]
    clear_examples(ws)
    n = write_rows(ws, [dict(item_code=l["item_code"], description=desc_of.get((contract_no, l["item_code"]), "(مش موجود في المقايسة)"),
                             site_code=l["site_code"], planned_qty=num(l["planned_qty"])) for l in lines])
    clear_examples(wb["مهمات أمر التكليف"])  # مفيش مهمات مخططة في الملفات الحالية
    save(wb, f"06_أمر_التكليف__{ao_key}.xlsx")
    print(f"     {n} سطر كمية، {len({l['site_code'] for l in lines if l['site_code']})} موقع")

# ═════════════ 07 كتالوج المهمات (الأصناف بس) ═════════════
wb = open_template("07_كتالوج_المهمات")
ws = wb["المهمات"]
clear_examples(ws)
group_name = {m["material_code"]: m["description"] for m in materials if m["is_group"] == "1"}
rows = []
for m in materials:
    if m["is_group"] == "1":
        continue
    top = m["material_code"].split(".")[0]
    rows.append(dict(material_code=m["material_code"], description=m["description"], unit=m["unit"],
                     supplied_by="المصرية للاتصالات",  # القائمة كلها "SUPPLY" — توريد المصرية للاتصالات
                     group=group_name.get(top, "") if top != m["material_code"] else ""))
n = write_rows(ws, rows)
save(wb, "07_كتالوج_المهمات.xlsx")
print(f"     {n} صنف")

# ═════════════ 11 المستخلصات السابقة ═════════════
wb = open_template("11_المستخلصات_السابقة")
ws = wb["المستخلصات"]
clear_examples(ws)
rows, unknown = [], 0
for inv in invoices:
    step, status = inv["mapped_step"], inv["mapped_status"]
    if step == "company" and status == "on_hold":
        step_ar = "بانتظار أمر التكليف"
    else:
        step_ar = STEP_AR.get(step, "")
    notes = [f"الموقف الأصلي: {inv['status_raw'] or '(فاضي)'}"]
    if not step_ar:
        unknown += 1
        notes.insert(0, "المحطة غير محددة — لازم تتحدد يدوي")
    if inv["region"]:
        notes.append(f"المنطقة: {inv['region']}")
    if inv["scope"]:
        notes.append(f"النطاق: {inv['scope']}")
    rows.append(dict(contract_no=inv["contract_no"], ao_no=inv["ao_no"], invoice_no=inv["invoice_no"], path_name=inv["path_name"],
                     amount_excl_vat=num(inv["amount_excl_vat"]), amount_incl_vat=num(inv["amount_incl_vat"]),
                     current_step=step_ar, holder_name=inv["holder_name"], notes=" | ".join(notes)))
n = write_rows(ws, rows, money_cols=("amount_excl_vat", "amount_incl_vat"))
for i, r in enumerate(rows):  # علّم الصفوف اللي محطتها ناقصة
    if not r["current_step"]:
        for j in range(1, ws.max_column + 1):
            ws.cell(row=3 + i, column=j).fill = WARN_FILL
save(wb, "11_المستخلصات_السابقة__بريدج.xlsx")
print(f"     {n} مستخلص ({unknown} محطتهم غير محددة)")
print("done")
